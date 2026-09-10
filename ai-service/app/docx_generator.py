"""Resume .docx rendering: built-in styles + optional DOCX templates from design_kb."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Twips

DESIGN_KB_DIR = Path(__file__).resolve().parent / "crews" / "design_kb"

BUILT_IN_TEMPLATES = ("classic", "modern", "compact")

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}", re.IGNORECASE)


def list_design_templates() -> list[dict]:
    """Built-ins plus any .docx/.pdf files dropped into design_kb/."""
    templates: list[dict] = [
        {
            "id": "classic",
            "kind": "builtin",
            "description": "Clean traditional headings, standard ATS-friendly layout",
        },
        {
            "id": "modern",
            "kind": "builtin",
            "description": "Stronger name hierarchy, accent emphasis, tighter spacing",
        },
        {
            "id": "compact",
            "kind": "builtin",
            "description": "Denser single-column layout for longer profiles",
        },
    ]
    if not DESIGN_KB_DIR.is_dir():
        return templates

    for path in sorted(DESIGN_KB_DIR.iterdir()):
        if path.name.startswith("."):
            continue
        suffix = path.suffix.lower()
        if suffix == ".docx":
            templates.append(
                {
                    "id": path.stem,
                    "kind": "docx_template",
                    "description": f"Local DOCX template file: {path.name}",
                    "path": str(path),
                }
            )
        elif suffix == ".pdf":
            templates.append(
                {
                    "id": path.stem,
                    "kind": "pdf_reference",
                    "description": (
                        f"PDF visual reference only ({path.name}); "
                        "if selected, render falls back to classic DOCX style"
                    ),
                    "path": str(path),
                }
            )
    return templates


def resolve_template(template_id: str | None) -> dict:
    templates = {t["id"]: t for t in list_design_templates()}
    if template_id and template_id in templates:
        return templates[template_id]
    return templates["classic"]


def generate_resume_docx(resume: dict, template_id: str | None = None) -> bytes:
    template = resolve_template(template_id)
    kind = template.get("kind", "builtin")

    if kind == "docx_template":
        filled = _try_fill_docx_template(resume, Path(template["path"]))
        if filled is not None:
            return filled
        # Template had no usable placeholders — fall back to classic styling.
        return _render_builtin(resume, "classic")

    style_id = template["id"] if kind == "builtin" else "classic"
    return _render_builtin(resume, style_id)


def _render_builtin(resume: dict, style_id: str) -> bytes:
    if style_id == "modern":
        return _render_modern(resume)
    if style_id == "compact":
        return _render_compact(resume)
    return _render_classic(resume)


def _render_classic(resume: dict) -> bytes:
    doc = Document()
    personal = resume.get("personal_info", {})
    doc.add_heading(personal.get("name") or "Resume", level=0)

    contact_bits = [v for v in (personal.get("email"), personal.get("phone"), personal.get("location")) if v]
    if contact_bits:
        doc.add_paragraph(" | ".join(contact_bits))

    _add_standard_sections(doc, resume, bullet_style="List Bullet", heading_level=1)
    return _save(doc)


def _render_modern(resume: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Twips(720)
    section.bottom_margin = Twips(720)
    section.left_margin = Twips(864)
    section.right_margin = Twips(864)

    personal = resume.get("personal_info", {})
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(personal.get("name") or "Resume")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    contact_bits = [v for v in (personal.get("email"), personal.get("phone"), personal.get("location")) if v]
    if contact_bits:
        contact = doc.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = contact.add_run(" · ".join(contact_bits))
        c_run.font.size = Pt(10)
        c_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _add_standard_sections(doc, resume, bullet_style="List Bullet", heading_level=1, modern_headings=True)
    return _save(doc)


def _render_compact(resume: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Twips(576)
    section.bottom_margin = Twips(576)
    section.left_margin = Twips(720)
    section.right_margin = Twips(720)

    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)

    personal = resume.get("personal_info", {})
    title = doc.add_paragraph()
    run = title.add_run(personal.get("name") or "Resume")
    run.bold = True
    run.font.size = Pt(16)

    contact_bits = [v for v in (personal.get("email"), personal.get("phone"), personal.get("location")) if v]
    if contact_bits:
        p = doc.add_paragraph(" | ".join(contact_bits))
        for r in p.runs:
            r.font.size = Pt(9)

    _add_standard_sections(doc, resume, bullet_style="List Bullet", heading_level=1, compact=True)
    return _save(doc)


def _add_standard_sections(
    doc: Document,
    resume: dict,
    *,
    bullet_style: str,
    heading_level: int,
    modern_headings: bool = False,
    compact: bool = False,
) -> None:
    def heading(text: str) -> None:
        h = doc.add_heading(text, level=heading_level)
        if modern_headings:
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        if compact:
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)

    if resume.get("professional_summary"):
        heading("Summary")
        doc.add_paragraph(resume["professional_summary"])

    if resume.get("work_experience"):
        heading("Experience")
        for entry in resume["work_experience"]:
            title_line = doc.add_paragraph()
            run = title_line.add_run(f"{entry.get('title', '')} — {entry.get('company', '')}")
            run.bold = True
            if compact:
                run.font.size = Pt(10)
            dates = " – ".join(v for v in (entry.get("start_date"), entry.get("end_date")) if v)
            if dates:
                date_run = title_line.add_run(f"  ({dates})")
                date_run.italic = True
                date_run.font.size = Pt(9 if compact else 10)
            for bullet in entry.get("bullets", []):
                bp = doc.add_paragraph(bullet, style=bullet_style)
                if compact:
                    bp.paragraph_format.space_after = Pt(0)

    if resume.get("education"):
        heading("Education")
        for entry in resume["education"]:
            line = " - ".join(
                v for v in (entry.get("degree"), entry.get("field"), entry.get("institution"), entry.get("year")) if v
            )
            doc.add_paragraph(line)

    if resume.get("skills"):
        heading("Skills")
        doc.add_paragraph(", ".join(resume["skills"]))

    if resume.get("projects"):
        heading("Projects")
        for entry in resume["projects"]:
            p = doc.add_paragraph()
            name_run = p.add_run(entry.get("name", ""))
            name_run.bold = True
            if entry.get("description"):
                doc.add_paragraph(entry["description"])

    if resume.get("languages"):
        heading("Languages")
        doc.add_paragraph(", ".join(resume["languages"]))


def _placeholder_values(resume: dict) -> dict[str, str]:
    personal = resume.get("personal_info") or {}
    experience_blocks: list[str] = []
    for entry in resume.get("work_experience") or []:
        header = f"{entry.get('title', '')} — {entry.get('company', '')}".strip(" —")
        dates = " – ".join(v for v in (entry.get("start_date"), entry.get("end_date")) if v)
        lines = [header + (f" ({dates})" if dates else "")]
        lines.extend(f"• {b}" for b in entry.get("bullets") or [])
        experience_blocks.append("\n".join(lines))

    education_lines = []
    for entry in resume.get("education") or []:
        education_lines.append(
            " - ".join(
                v for v in (entry.get("degree"), entry.get("field"), entry.get("institution"), entry.get("year")) if v
            )
        )

    project_blocks = []
    for entry in resume.get("projects") or []:
        block = entry.get("name", "")
        if entry.get("description"):
            block = f"{block}\n{entry['description']}"
        project_blocks.append(block)

    return {
        "name": personal.get("name") or "",
        "email": personal.get("email") or "",
        "phone": personal.get("phone") or "",
        "location": personal.get("location") or "",
        "summary": resume.get("professional_summary") or "",
        "skills": ", ".join(resume.get("skills") or []),
        "languages": ", ".join(resume.get("languages") or []),
        "experience": "\n\n".join(experience_blocks),
        "education": "\n".join(education_lines),
        "projects": "\n\n".join(project_blocks),
    }


def _try_fill_docx_template(resume: dict, path: Path) -> bytes | None:
    try:
        doc = Document(str(path))
    except Exception:
        return None

    values = _placeholder_values(resume)
    replaced = 0

    def subst(text: str) -> str:
        nonlocal replaced

        def _repl(match: re.Match[str]) -> str:
            nonlocal replaced
            key = match.group(1).lower()
            if key in values:
                replaced += 1
                return values[key]
            return match.group(0)

        return _PLACEHOLDER_RE.sub(_repl, text)

    for paragraph in doc.paragraphs:
        if _PLACEHOLDER_RE.search(paragraph.text):
            # Replace at paragraph level to keep it simple across runs.
            new_text = subst(paragraph.text)
            if new_text != paragraph.text:
                for run in paragraph.runs:
                    run.text = ""
                if paragraph.runs:
                    paragraph.runs[0].text = new_text
                else:
                    paragraph.add_run(new_text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _PLACEHOLDER_RE.search(paragraph.text):
                        new_text = subst(paragraph.text)
                        if new_text != paragraph.text:
                            for run in paragraph.runs:
                                run.text = ""
                            if paragraph.runs:
                                paragraph.runs[0].text = new_text
                            else:
                                paragraph.add_run(new_text)

    if replaced == 0:
        return None
    return _save(doc)


def _save(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
